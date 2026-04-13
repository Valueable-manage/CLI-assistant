"""
文件解析器 - 把各种格式的文件提取为纯文本
"""
import csv
import io


def parse_file(filename: str, content: bytes) -> str:
    """根据文件扩展名选择对应的解析方式"""
    ext = filename.rsplit(".", 1)[-1].lower()

    if ext == "txt" or ext == "md":
        return _parse_text(content)
    elif ext == "csv":
        return _parse_csv(content)
    elif ext == "docx":
        return _parse_docx(content)
    elif ext in ("xlsx", "xls"):
        return _parse_excel(content, ext)
    elif ext == "pdf":
        return _parse_pdf(content)
    else:
        raise ValueError(f"不支持的文件格式: .{ext}（支持 txt/md/csv/docx/xlsx/pdf）")


def _parse_text(content: bytes) -> str:
    return content.decode("utf-8", errors="ignore")


def _parse_csv(content: bytes) -> str:
    text = content.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    rows = ["\t".join(row) for row in reader]
    return "\n".join(rows)


def _parse_docx(content: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def _parse_excel(content: bytes, ext: str) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        parts.append(f"【工作表：{sheet}】")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _parse_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"【第{i+1}页】\n{text}")
    return "\n\n".join(pages)
